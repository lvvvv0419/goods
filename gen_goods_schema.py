# -*- coding: utf-8 -*-
"""根据 01(2).xlsx 生成 MySQL `goods` 数据库建表/导入 SQL。"""
import openpyxl

XLSX = r"C:\Users\user\Desktop\01(2).xlsx"

# 列字母 -> 英文列名（按 Excel 列顺序，已排除空列 DJ 和末列 EC「物料二维码」）
EN = {
    # 基础物料信息 A-R
    "A": "seq", "B": "code", "C": "material_name", "D": "spec_model",
    "E": "mnemonic_code", "F": "material_attr", "G": "config_production",
    "H": "inventory_category", "I": "allow_purchase", "J": "allow_sales",
    "K": "allow_inventory", "L": "allow_production", "M": "allow_outsourcing",
    "N": "allow_asset", "O": "tax_classification", "P": "default_tax_rate",
    "Q": "base_unit", "R": "is_kit",
    # 物料销售信息 S-AD
    "S": "sales_unit", "T": "sales_pricing_unit", "U": "order_batch",
    "V": "min_batch", "W": "max_batch", "X": "over_delivery_upper_limit",
    "Y": "over_delivery_lower_limit", "Z": "allow_return", "AA": "return_inspection",
    "AB": "agency_sales_discount_ratio", "AC": "allow_publish_to_order_platform",
    "AD": "product_category",
    # 物料采购信息 AE-BB
    "AE": "purchase_unit", "AF": "purchase_pricing_unit", "AG": "purchase_group",
    "AH": "purchaser", "AI": "default_supplier", "AJ": "supplier_qualification",
    "AK": "source_control", "AL": "need_purchase_request", "AM": "purchase_batch",
    "AN": "purchase_min_order_qty", "AO": "receiving_upper_limit_pct",
    "AP": "receiving_lower_limit_pct", "AQ": "delivery_early_days",
    "AR": "delivery_late_days", "AS": "quota_management", "AT": "quota_sequence",
    "AU": "min_split_qty", "AV": "agency_purchase_markup_ratio",
    "AW": "is_vmi_business", "AX": "outsourcing_unit", "AY": "outsourcing_pricing_unit",
    "AZ": "outsourcing_daily_output", "BA": "outsourcing_default_line",
    "BB": "outsourcing_enable_daily_schedule",
    # 物料制造信息 BC-BX
    "BC": "production_workshop", "BD": "production_unit", "BE": "receipt_over_ratio",
    "BF": "receipt_under_ratio", "BG": "serial_carried_to_parent",
    "BH": "child_item_unit", "BI": "can_be_coproduct", "BJ": "default_routing",
    "BK": "standard_work_hours_sec", "BL": "issue_method", "BM": "backflush_timing",
    "BN": "issue_warehouse", "BO": "issue_location", "BP": "over_issue_control",
    "BQ": "min_issue_batch", "BR": "is_critical_part", "BS": "is_kit_part",
    "BT": "product_model", "BU": "model_material", "BV": "production_daily_output",
    "BW": "production_default_line", "BX": "production_enable_daily_schedule",
    # 物料MRP运算 BY-DI
    "BY": "planning_strategy", "BZ": "manufacturing_strategy",
    "CA": "ordering_strategy", "CB": "planning_area", "CC": "ordering_interval",
    "CD": "fixed_lead_time", "CE": "variable_lead_time",
    "CF": "variable_lead_time_batch", "CG": "inspection_lead_time",
    "CH": "max_order_qty", "CI": "mrp_min_order_qty", "CJ": "batch_increment",
    "CK": "fixed_economic_batch", "CL": "split_batch",
    "CM": "batch_split_interval_days", "CN": "demand_time_fence",
    "CO": "planning_time_fence", "CP": "planning_group", "CQ": "planner",
    "CR": "planning_flag", "CS": "match_supply_demand_by_batch",
    "CT": "mrp_safety_stock", "CU": "reservation_type", "CV": "allow_manual_reservation",
    "CW": "ato_consumption_plan", "CX": "product_series", "CY": "consumption_qty",
    "CZ": "allow_early_days", "DA": "allow_partial_early", "DB": "early_grace_days",
    "DC": "allow_late_days", "DD": "allow_partial_late", "DE": "late_grace_days",
    "DF": "supply_source", "DG": "order_progress_group", "DH": "mrp_merge_demand",
    "DI": "mrp_merge_by_order",
    # 物料仓储信息 DK-EB
    "DK": "inventory_unit", "DL": "auxiliary_unit", "DM": "warehouse",
    "DN": "location", "DO": "enable_batch_management", "DP": "enable_shelf_life_control",
    "DQ": "batch_auxiliary_info", "DR": "shelf_life_unit", "DS": "shelf_life",
    "DT": "enable_serial_management", "DU": "serial_number_rule", "DV": "business_scope",
    "DW": "serial_generation_timing", "DX": "min_inventory", "DY": "max_inventory",
    "DZ": "warehouse_safety_stock", "EA": "reorder_point",
    "EB": "standard_qty_per_box",
}

# 是/否 布尔列 -> TINYINT(1)
BOOL = {"I", "J", "K", "L", "M", "N", "R", "Z", "AA", "AC", "AK", "AL", "AS",
        "AW", "BB", "BG", "BI", "BR", "BS", "BX", "CS", "CV", "DA", "DD", "DH",
        "DI", "DO", "DP", "DT"}

# 税率/比例/提前期 -> DECIMAL
DEC = {"P", "AB", "AO", "AP", "AV", "BE", "BF", "CD", "CE", "CG"}

# 长文本 -> VARCHAR(255)
LONG = {"C", "D", "AI", "AJ", "BJ", "CW", "DF", "DG", "DQ", "DU", "DV"}

REQUIRED = {"B", "C", "F", "Q"}  # 带 * 的必填字段


def col_type(letter, val):
    if letter in BOOL:
        return "TINYINT(1)"
    if letter in DEC:
        return "DECIMAL(12,4)"
    if isinstance(val, bool):
        return "TINYINT(1)"
    if isinstance(val, int):
        return "INT"
    if isinstance(val, float):
        return "DECIMAL(12,4)"
    return f"VARCHAR({255 if letter in LONG else 64})"


def main():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    ws = wb.active
    headers = {c.column: c.value for c in ws[2] if c.value is not None}
    data = {c.column: c.value for c in ws[3]}

    from openpyxl.utils import get_column_letter
    cols = []
    for col_idx in sorted(headers):
        letter = get_column_letter(col_idx)
        if letter == "EC" or letter not in EN:
            continue
        cn = headers[col_idx].rstrip("*").strip()
        en = EN[letter]
        v = data.get(col_idx)
        t = col_type(letter, v)
        nn = " NOT NULL" if letter in REQUIRED else " NULL"
        cols.append((letter, en, t, cn, nn, v))

    lines = []
    lines.append("SET NAMES utf8mb4;")
    lines.append("DROP DATABASE IF EXISTS `goods`;")
    lines.append("CREATE DATABASE `goods` DEFAULT CHARACTER SET utf8mb4 COLLATE utf8mb4_0900_ai_ci;")
    lines.append("USE `goods`;")
    lines.append("")
    lines.append("CREATE TABLE `material` (")
    lines.append("  `id` BIGINT UNSIGNED NOT NULL AUTO_INCREMENT COMMENT '主键ID',")
    for letter, en, t, cn, nn, v in cols:
        lines.append(f"  `{en}` {t}{nn} COMMENT '{cn}',")
    lines.append("  PRIMARY KEY (`id`),")
    lines.append("  UNIQUE KEY `uk_code` (`code`)")
    lines.append(") ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COMMENT='物料基础信息';")
    lines.append("")

    # INSERT 示例数据
    names = [en for _, en, _, _, _, _ in cols]
    vals = []
    for letter, en, t, cn, nn, v in cols:
        if v is None:
            vals.append("NULL")
        elif letter in BOOL:
            vals.append("1" if v == "是" else "0")
        elif isinstance(v, (int, float)):
            vals.append(repr(v))
        else:
            s = str(v).replace("\\", "\\\\").replace("'", "''")
            vals.append(f"'{s}'")
    lines.append("INSERT INTO `material` (`" + "`, `".join(names) + "`) VALUES")
    lines.append("  (" + ", ".join(vals) + ");")
    lines.append("")

    sql = "\n".join(lines)
    with open("goods.sql", "w", encoding="utf-8") as f:
        f.write(sql)
    print(f"共 {len(cols)} 个字段，SQL 已写入 goods.sql")
    print("=" * 60)
    print(sql)


if __name__ == "__main__":
    main()
